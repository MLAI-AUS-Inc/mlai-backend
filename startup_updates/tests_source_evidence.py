"""Source coverage regressions; these tests do not construct a database."""
import io
import unittest
from datetime import datetime, timezone
from types import SimpleNamespace
from unittest.mock import patch

from startup_updates.source_evidence import period_gmail_bundle, source_receipt
from startup_updates.manual_documents import parse_manual_document


class SourceEvidenceTests(unittest.TestCase):
    def test_february_thread_survives_later_replies_without_future_claims(self):
        messages = [
            {"message_id": "feb", "internal_date": "2026-02-12T12:00:00+00:00", "cleaned_text": "February pilot signed."},
            {"message_id": "sep", "internal_date": "2026-09-05T12:00:00+00:00", "cleaned_text": "September pilot cancelled."},
        ]
        thread = SimpleNamespace(gmail_thread_id="thread", message_payloads=messages)
        bundle = period_gmail_bundle(thread, start=datetime(2026, 2, 1, tzinfo=timezone.utc), end=datetime(2026, 2, 28, 23, 59, 59, tzinfo=timezone.utc), attachments=[])
        self.assertEqual(bundle["source_message_ids"], ["feb"])
        self.assertIn("pilot signed", bundle["cleaned_text"])
        self.assertNotIn("cancelled", bundle["cleaned_text"])
        self.assertEqual(bundle["omitted_message_count"], 0)

    def test_receipts_change_with_period_source_or_extractor_version(self):
        run = SimpleNamespace(run_request={"draft_months": ["2026-02-01"]})
        bundle = {"text": "Original February result"}
        before = source_receipt(run, "notion", "page", bundle)["fingerprint"]
        self.assertEqual(before, source_receipt(run, "notion", "page", bundle)["fingerprint"])
        self.assertNotEqual(before, source_receipt(run, "notion", "page", {"text": "Edited result"})["fingerprint"])
        run.run_request["draft_months"] = ["2026-08-01"]
        self.assertNotEqual(before, source_receipt(run, "notion", "page", bundle)["fingerprint"])
        run.run_request["draft_months"] = ["2026-02-01"]
        with patch("startup_updates.source_evidence.EXTRACTION_VERSION", "future-parser"):
            self.assertNotEqual(before, source_receipt(run, "notion", "page", bundle)["fingerprint"])

    def test_parser_cache_does_not_change_the_underlying_source_fingerprint(self):
        run = SimpleNamespace(run_request={})
        bundle = {"attachments": [{"id": 1, "raw_content_base64": "YWJj", "extracted_text": "", "parse_notes": ""}]}
        before = source_receipt(run, "gmail", "thread", bundle)
        bundle["attachments"][0].update(extracted_text="Table fact", parse_notes="docx_parsed", extraction_status="processed")
        after = source_receipt(run, "gmail", "thread", bundle)
        self.assertEqual(before["fingerprint"], after["fingerprint"])
        self.assertNotIn("raw_content_base64", after["bundle"]["attachments"][0])
        self.assertIn("binary_hash", after["bundle"]["attachments"][0])

    def test_uploaded_docx_table_after_old_character_limit_is_extracted(self):
        from docx import Document
        document = Document()
        document.add_paragraph("Background. " * 2500)
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "2026-08-23 milestone"
        table.rows[0].cells[1].text = "First accessible offline release shipped."
        stream = io.BytesIO()
        document.save(stream)
        parsed = parse_manual_document(filename="evidence.docx", content_type="", raw_bytes=stream.getvalue())
        self.assertEqual(parsed.extraction_status, "processed")
        self.assertGreater(parsed.extracted_text.index("First accessible offline release"), 12000)
        self.assertIn("2026-08-23 milestone | First accessible offline release shipped.", parsed.extracted_text)

    def test_spreadsheet_blank_column_does_not_shift_its_metric_label(self):
        parsed = parse_manual_document(filename="evidence.csv", content_type="text/csv", raw_bytes=b"Month,Revenue,Customers\nAugust,,7")
        self.assertIn("August | [empty] | 7", parsed.extracted_text)


class NotionCoverageTests(unittest.TestCase):
    def test_nested_blocks_and_fact_beyond_eighty_blocks_are_kept(self):
        from startup_updates.api_views import _build_notion_page_bundle
        blocks = [{"id": str(i), "type": "paragraph", "paragraph": {"rich_text": [{"plain_text": f"Paragraph {i}"}]}} for i in range(101)]
        blocks[0]["has_children"] = True
        nested = {"id": "nested", "type": "table_row", "table_row": {"cells": [[{"plain_text": "February 2026"}], [{"plain_text": "Retention review completed"}]]}}
        with patch("startup_updates.api_views._fetch_notion_children", side_effect=lambda connection, block_id: blocks if block_id == "page" else [nested]):
            bundle = _build_notion_page_bundle(None, {"id": "page", "last_edited_time": "2026-09-01T00:00:00Z"})
        self.assertIn("Paragraph 100", bundle["cleaned_text"])
        self.assertIn("February 2026 | Retention review completed", bundle["cleaned_text"])
        self.assertEqual(len(bundle["source_block_ids"]), 102)
        self.assertEqual(bundle["omitted_block_count"], 0)
        self.assertEqual(bundle["last_edited_time"], "2026-09-01T00:00:00Z")

    def test_nonadvancing_notion_cursor_fails_instead_of_silent_truncation(self):
        from startup_updates.api_views import _fetch_notion_children
        response = SimpleNamespace(raise_for_status=lambda: None, json=lambda: {"results": [], "has_more": True, "next_cursor": "stuck"})
        with patch("startup_updates.api_views.requests.get", return_value=response), patch("startup_updates.api_views._notion_headers", return_value={}):
            with self.assertRaisesRegex(ValueError, "did not advance"):
                _fetch_notion_children(None, "page")

class AccountingReportTests(unittest.TestCase):
    def test_balance_sheet_uses_current_column_and_ignores_comparative_year(self):
        from decimal import Decimal
        from startup_updates.services import _parse_xero_balance_sheet_report, _xero_balance_sheet_date
        from datetime import date
        payload = {"Reports": [{"Rows": [
            {"RowType": "Header", "Cells": [{"Value": ""}, {"Value": "31 Aug 2026"}, {"Value": "31 Aug 2025"}]},
            {"RowType": "SummaryRow", "Cells": [{"Value": "Total Bank"}, {"Value": "1200.50"}, {"Value": "9000.00"}]},
        ]}]}
        self.assertEqual(_parse_xero_balance_sheet_report(payload)["cash"]["amount"], Decimal("1200.50"))
        self.assertEqual(_xero_balance_sheet_date(payload), date(2026, 8, 31))
        payload["Reports"][0]["Rows"][1]["Cells"][1]["Value"] = ""
        self.assertIsNone(_parse_xero_balance_sheet_report(payload)["cash"])
        payload["Reports"][0]["Rows"][1]["Cells"][1]["Value"] = "0"
        self.assertEqual(_parse_xero_balance_sheet_report(payload)["cash"]["amount"], 0)

    def test_profit_and_loss_uses_requested_column_even_when_comparison_is_larger(self):
        from startup_updates.services import _parse_xero_profit_and_loss_report
        payload = self.report([("Total Income", 100)])
        payload["Reports"][0]["Rows"][0]["Rows"][0]["Cells"].append({"Value": "900"})
        self.assertEqual(_parse_xero_profit_and_loss_report(payload)["revenue"]["amount"], 100)

    def report(self, rows):
        return {"Reports": [{"Rows": [{"RowType": "Section", "Title": "Income", "Rows": [
            {"RowType": "SummaryRow", "Cells": [{"Value": label}, {"Value": str(amount)}]} for label, amount in rows
        ]}]}]}

    def test_report_revenue_uses_total_and_keeps_negative_adjustments(self):
        from startup_updates.services import _parse_xero_profit_and_loss_report
        parsed = _parse_xero_profit_and_loss_report(self.report([("Other Income", 500), ("Total Income", -50), ("Total Operating Expenses", -20)]))
        self.assertEqual(parsed["revenue"]["amount"], -50)
        self.assertEqual(parsed["monthly_costs"]["amount"], -20)

    def test_single_income_category_is_not_whole_business_revenue(self):
        from startup_updates.services import _parse_xero_profit_and_loss_report
        parsed = _parse_xero_profit_and_loss_report(self.report([("Other Income", 500)]))
        self.assertIsNone(parsed["revenue"])

    def test_cost_total_cannot_be_replaced_by_individual_expense_account(self):
        from startup_updates.services import _parse_xero_profit_and_loss_report
        payload = self.report([("Total Income", 1000), ("Contractor Expenses", 300), ("General Expenses", 200), ("Total Operating Expenses", 500), ("Net Profit", 500)])
        parsed = _parse_xero_profit_and_loss_report(payload)
        self.assertEqual(parsed["monthly_costs"]["amount"], 500)
        self.assertEqual(parsed["revenue"]["amount"] - parsed["monthly_costs"]["amount"], parsed["net"]["amount"])

    def test_summary_cost_of_sales_precedes_individual_cost_account(self):
        from startup_updates.services import _parse_xero_profit_and_loss_report
        parsed = _parse_xero_profit_and_loss_report(self.report([("Cost of Goods Sold", 10), ("Total Cost of Sales", 30), ("Total Operating Expenses", 20)]))
        self.assertEqual(parsed["monthly_costs"]["amount"], 50)

    def test_missing_expense_total_stays_unknown(self):
        from startup_updates.services import _parse_xero_profit_and_loss_report
        parsed = _parse_xero_profit_and_loss_report(self.report([("Contractor Expenses", 300), ("General Expenses", 200)]))
        self.assertIsNone(parsed["monthly_costs"])

    def test_report_breakdown_keeps_signed_adjustments(self):
        from startup_updates.services import _xero_report_breakdown_rows
        from decimal import Decimal
        rows = _xero_report_breakdown_rows([{"label": "Ticket adjustment", "normalized_label": "ticket adjustment", "amount": Decimal("-20"), "section": "Income", "normalized_section": "income"}], section_terms=("income",))
        self.assertEqual(rows[0]["amount"], "-20")

class ClassificationContractTests(unittest.TestCase):
    def test_cache_changes_with_period_or_startup_context_but_not_run_id(self):
        from startup_updates.source_evidence import classification_version
        run = SimpleNamespace(run_id="first", run_request={"draft_months": ["2026-02-01"], "startup_context": {"company_name": "Original"}})
        original = classification_version(run)
        run.run_id = "retry"
        self.assertEqual(classification_version(run), original)
        run.run_request["startup_context"]["company_name"] = "Renamed"
        self.assertNotEqual(classification_version(run), original)
        renamed = classification_version(run)
        run.run_request["draft_months"] = ["2026-08-01"]
        self.assertNotEqual(classification_version(run), renamed)


class ChartEvidenceTests(unittest.TestCase):
    def metric(self, **changes):
        from datetime import date
        from decimal import Decimal
        from integrations.tests_connectors import _xero_profit_and_loss_report
        from startup_updates.evidence_contract import content_hash
        report = _xero_profit_and_loss_report(total_income='100', total_expenses='70', net_profit='30')
        metadata = {'source_metric': 'xero_profit_and_loss_monthly_costs', 'report_hash': content_hash(report),
            'report_payload': report, 'accounting_basis': 'accrual', 'report_start_date': '2026-08-01', 'report_end_date': '2026-08-31'}
        metric = SimpleNamespace(source_provider='xero', metric_key='monthlyCosts', period_month=date(2026,8,1), value_number=Decimal('70'), source_metadata=metadata)
        for key,value in changes.items(): setattr(metric,key,value)
        return metric

    def test_historical_cost_must_match_corrected_report_parser(self):
        from decimal import Decimal
        from startup_updates.services import _verified_chart_observation
        self.assertTrue(_verified_chart_observation(self.metric()))
        # An old contractor-only total remains invalid even with the same raw report.
        self.assertFalse(_verified_chart_observation(self.metric(value_number=Decimal('40'))))
        self.assertFalse(_verified_chart_observation(self.metric(source_metadata={})))

    def test_wrong_hash_period_basis_and_source_label_are_not_chart_evidence(self):
        from startup_updates.services import _verified_chart_observation
        for key,value in [('report_hash','changed'),('report_start_date','2026-07-01'),('report_end_date','2026-09-01'),('report_end_date',None),('accounting_basis','cash'),('source_metric','contractor_expenses')]:
            with self.subTest(key=key,value=value):
                metric=self.metric();metric.source_metadata[key]=value
                self.assertFalse(_verified_chart_observation(metric))

    def test_stripe_requires_versioned_paid_sales_and_never_supplies_costs(self):
        from startup_updates.services import _verified_chart_observation
        metadata={'definition_version':2,'basis':'paid_stripe_invoice_sales_excluding_tax'}
        metric=self.metric(source_provider='financial',metric_key='revenue',source_metadata=metadata)
        self.assertTrue(_verified_chart_observation(metric))
        metric.metric_key='monthlyCosts';self.assertFalse(_verified_chart_observation(metric))
        metric.metric_key='revenue';metadata.pop('definition_version');self.assertFalse(_verified_chart_observation(metric))


RUN_REQUEST = {
    'start_date': '2026-10-01', 'end_date': '2026-10-18',
    'backfill_window_start': '2026-10-01T00:00:00+10:00',
    'backfill_window_end': '2026-10-18T12:00:00+11:00',
    'draft_months': ['2026-10-01'], 'reporting_timezone': 'Australia/Melbourne',
    'startup_context': {'company_name': 'Synthetic 🦘', 'notes': 'Founder notes with 漢字.'},
}
BUNDLE = {'cleaned_text': 'The pilot shipped.', 'message_payloads': [{'message_id': 'message-1', 'cleaned_text': 'The pilot shipped.'}],
    'attachments': [{'id': 7, 'filename': 'report.pdf', 'raw_content_base64': 'c3ludGhldGlj',
        'extracted_text': 'Pilot evidence.', 'extraction_status': 'processed', 'parse_notes': '', 'last_error': '', 'metadata': {'parser': 1}}]}


class HoistedStartupContextTests(unittest.TestCase):
    def test_golden_fingerprint_retains_main_contract(self):
        import copy
        from startup_updates.evidence_contract import content_hash
        original = copy.deepcopy(BUNDLE)
        run = SimpleNamespace(run_request=RUN_REQUEST)
        receipt = source_receipt(run, "gmail", "thread-1", BUNDLE)
        # Computed on origin/main 6bbe177f before changing source_receipt.
        self.assertEqual(receipt["fingerprint"], "82d8b38cdfe55d624e5c5709b755be805158325c790ce4af0f29ea41970345bd")
        self.assertNotIn("startup_context", receipt)
        self.assertEqual(receipt["startup_context_hash"], content_hash(RUN_REQUEST["startup_context"]))
        self.assertEqual(BUNDLE, original)
        self.assertEqual(run.run_request, RUN_REQUEST)

    def test_legacy_context_is_hoisted_without_mutating_receipts(self):
        import copy
        from startup_updates.source_evidence import hoist_startup_context
        from startup_updates.evidence_contract import content_hash
        receipts = {"gmail:1": {"startup_context": {"notes": "Frozen 🦘"}, "fingerprint": "one"},
            "gmail:2": {"startup_context": {"notes": "Frozen 🦘"}, "fingerprint": "two"}}
        original = copy.deepcopy(receipts)
        normalized, context = hoist_startup_context(receipts, fallback={"notes": "Live profile"})
        self.assertEqual(context, {"notes": "Frozen 🦘"})
        self.assertEqual(receipts, original)
        for key, receipt in normalized.items():
            self.assertNotIn("startup_context", receipt)
            self.assertEqual(receipt["startup_context_hash"], content_hash(context))
            self.assertEqual(receipt["fingerprint"], original[key]["fingerprint"])

    def test_distinct_legacy_contexts_are_retained_in_full(self):
        from startup_updates.source_evidence import hoist_startup_context
        receipts = {"one": {"startup_context": {"notes": "First"}},
            "two": {"startup_context": {"notes": "Different"}}}
        normalized, context = hoist_startup_context(receipts, fallback={"notes": "Live"})
        self.assertEqual(normalized, receipts)
        self.assertIsNone(context)

    def test_modern_receipts_use_frozen_run_context_fallback(self):
        from startup_updates.source_evidence import hoist_startup_context
        context = {"notes": "Frozen"}
        receipt = source_receipt(SimpleNamespace(run_request={"startup_context": context}), "gmail", "1", {})
        normalized, hoisted = hoist_startup_context({"gmail:1": receipt}, fallback=context)
        self.assertEqual(normalized, {"gmail:1": receipt})
        self.assertEqual(hoisted, context)
        hoisted["notes"] = "Mutated"
        self.assertEqual(context["notes"], "Frozen")
